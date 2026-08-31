"""Document parsing: stored bytes -> canonical text + structure + page map.

Canonical text is the "\\n\\n" join of block texts; every downstream citation
is a char offset into it, so the parsers own that invariant (PROGRESS W3).
DOCX clause structure comes from paragraph heading styles; PDF lines get a
conservative clause-number regex, since pdfplumber has no style information.
Plain text never yields headings (ADR-005: no clause-level chunks for it).
Scanned PDFs/OCR are out of scope (PROGRESS W3); a text-less PDF parses to
an empty document.

Pure functions on bytes — blocking work is the caller's concern
(asyncio.to_thread in the ingestion service).
"""

import enum
import io
import re
from collections.abc import Sequence
from dataclasses import dataclass

import pdfplumber
from docx import Document as DocxDocument

_PDF_MIME = "application/pdf"
_TEXT_MIME = "text/plain"
_DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


class BlockKind(enum.Enum):
    heading = "heading"
    paragraph = "paragraph"


@dataclass(frozen=True)
class ParsedBlock:
    kind: BlockKind
    text: str
    start: int
    end: int
    page: int | None


@dataclass(frozen=True)
class ParsedDocument:
    text: str
    blocks: list[ParsedBlock]
    # [{"page": n, "start": s, "end": e}, ...] ascending, non-overlapping;
    # separators between pages belong to the earlier page's span.
    page_map: list[dict[str, int]]


class UnsupportedParseError(Exception):
    def __init__(self, mime_type: str) -> None:
        self.mime_type = mime_type
        super().__init__(f"no parser for media type {mime_type!r}")


# "1.", "1.2", "8.2.1)", "Section 3" followed by title text. A bare integer
# followed by words ("30 days after delivery …") is body prose, not a clause
# boundary — it matches none of the alternatives (no internal dot, no
# trailing dot/paren, no Article/Section/Clause prefix).
_HEADING_PATTERN = re.compile(
    r"^(?:(?:article|section|clause)\s+\d+(?:\.\d+)*[.)]?|\d+(?:\.\d+)+[.)]?|\d+\.)\s+\S",
    re.IGNORECASE,
)


def parse_pdf(content: bytes) -> ParsedDocument:
    entries: list[tuple[BlockKind, str, int | None]] = []
    with pdfplumber.open(io.BytesIO(content)) as pdf:
        for number, page in enumerate(pdf.pages, start=1):
            for line in (page.extract_text() or "").splitlines():
                line = line.strip()
                if not line:
                    continue
                kind = BlockKind.heading if _HEADING_PATTERN.match(line) else BlockKind.paragraph
                entries.append((kind, line, number))
    return _build(entries)


def parse_docx(content: bytes) -> ParsedDocument:
    document = DocxDocument(io.BytesIO(content))
    entries: list[tuple[BlockKind, str, int | None]] = []
    for para in document.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        style = para.style.name if para.style is not None else ""
        kind = BlockKind.heading if style.startswith(("Heading", "Title")) else BlockKind.paragraph
        entries.append((kind, text, None))
    return _build(entries)


def parse_text(content: bytes) -> ParsedDocument:
    text = content.decode("utf-8")
    entries = [
        (BlockKind.paragraph, part.strip(), None)
        for part in re.split(r"\n\s*\n", text)
        if part.strip()
    ]
    return _build(entries)


_PARSERS = {
    _PDF_MIME: parse_pdf,
    _DOCX_MIME: parse_docx,
    _TEXT_MIME: parse_text,
}


def parse(mime_type: str, content: bytes) -> ParsedDocument:
    parser = _PARSERS.get(mime_type)
    if parser is None:
        raise UnsupportedParseError(mime_type)
    return parser(content)


def _build(entries: Sequence[tuple[BlockKind, str, int | None]]) -> ParsedDocument:
    """Join block texts with "\\n\\n" and compute offsets + page spans in one pass."""
    parts: list[str] = []
    blocks: list[ParsedBlock] = []
    page_map: list[dict[str, int]] = []
    pos = 0
    for kind, text, page in entries:
        if parts:
            parts.append("\n\n")
            pos += 2
        start = pos
        pos += len(text)
        parts.append(text)
        blocks.append(ParsedBlock(kind=kind, text=text, start=start, end=pos, page=page))
        if page is not None and (not page_map or page_map[-1]["page"] != page):
            if page_map:
                page_map[-1]["end"] = start
            page_map.append({"page": page, "start": start})
    if page_map:
        page_map[-1]["end"] = pos
    return ParsedDocument(text="".join(parts), blocks=blocks, page_map=page_map)
