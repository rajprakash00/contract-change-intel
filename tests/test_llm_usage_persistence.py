"""Integration tests for per-call LLM usage persistence (issue #30): job
runs and the search surface persist one usage row per LLM call — matching
the client's one-log-line-per-call contract — and the aggregate reads
group spend per tenant, job kind, and job.

Real Postgres; the LLM wire is faked at the httpx2 transport seam
(tests/fake_openai), per the AGENTS.md exception.
"""

import io
import json
import uuid
from collections.abc import AsyncIterator
from contextlib import asynccontextmanager
from pathlib import Path

import pytest
from docx import Document as DocxDocument
from httpx import AsyncClient
from sqlalchemy import delete, select
from sqlalchemy.ext.asyncio import AsyncSession

import app.db as db
import app.repositories.document_chunks as chunks_repo
import app.repositories.document_texts as texts_repo
import app.repositories.documents as documents_repo
import app.repositories.llm_usage as usage_repo
from app.api.deps import get_llm_client
from app.config import get_settings
from app.llm.client import LlmUsage, OpenAiClient
from app.main import app
from app.models.change_report_job import ChangeReportJob
from app.models.document import Document, DocumentStatus
from app.models.document_chunk import EMBEDDING_DIMENSIONS, DocumentChunk
from app.models.document_text import DocumentText
from app.models.extraction_job import ExtractionJob, ExtractionJobStatus
from app.models.ingestion_job import IngestionJob
from app.models.llm_usage import LlmUsageRecord, UsageJobKind
from app.models.review_item import ReviewItem
from app.services.change_report import enqueue_change_report, run_next_change_report_job
from app.services.extraction import enqueue_extraction, run_next_extraction_job
from app.services.ingestion import enqueue_ingestion, run_next_ingestion_job
from app.storage import local
from tests.fake_jwks import bearer
from tests.fake_openai import (
    fake_embedding_client,
    fake_llm_client,
    fake_llm_client_queue,
    fake_llm_with_embeddings_client,
)

BASE_TEXT = "2.1 Delivery\n\nLICENSOR shall deliver within 14 days."
AMENDED_TEXT = "2.1 Delivery\n\nLICENSOR shall deliver within 30 days."
EXPLANATION_OUTPUT = json.dumps(
    {"changes": [{"index": 0, "description": "The delivery window doubles.", "severity": "high"}]}
)
IMPACT_OUTPUT = json.dumps({"impacts": [{"index": 0, "confidence": 0.8}]})
CHUNK_EMBEDDING = [1.0] + [0.0] * (EMBEDDING_DIMENSIONS - 1)

EXTRACTION_OUTPUT = json.dumps(
    {
        "obligations": [
            {
                "clause_ref": "8.2",
                "description": "Supplier shall deliver monthly status reports",
                "owner": "Supplier",
                "citation": {"char_start": 0, "char_end": 8},
                "confidence": 0.9,
            }
        ],
        "defined_terms": [],
    }
)

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def contract_docx() -> bytes:
    """A two-clause DOCX; each clause becomes one chunk (ADR-005)."""
    buffer = io.BytesIO()
    doc = DocxDocument()
    doc.add_heading("Section 1: Payment", level=1)
    doc.add_paragraph("Supplier shall pay within thirty days of invoice.")
    doc.add_heading("Section 2: Delivery", level=1)
    doc.add_paragraph("Goods shall be delivered to the customer site.")
    doc.save(buffer)
    return buffer.getvalue()


@asynccontextmanager
async def session() -> AsyncIterator[AsyncSession]:
    async with db.get_sessionmaker()() as s:
        yield s


@pytest.fixture(autouse=True)
async def engine() -> AsyncIterator[None]:
    db.init_engine(get_settings().database_url)
    yield
    await db.dispose_engine()


@pytest.fixture(autouse=True)
def data_dir(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> Path:
    monkeypatch.setenv("DATA_DIR", str(tmp_path))
    return tmp_path


@pytest.fixture(autouse=True)
async def clean_tables(data_dir: Path) -> AsyncIterator[None]:
    yield
    async with session() as s:
        await s.execute(delete(LlmUsageRecord))
        await s.execute(delete(ReviewItem))
        await s.execute(delete(ChangeReportJob))
        await s.execute(delete(ExtractionJob))
        await s.execute(delete(IngestionJob))
        await s.execute(delete(DocumentChunk))
        await s.execute(delete(DocumentText))
        await s.execute(delete(Document))
        await s.commit()
    app.dependency_overrides.pop(get_llm_client, None)


def usage(
    *,
    tenant_id: uuid.UUID,
    job_type: UsageJobKind,
    job_id: uuid.UUID | None = None,
    prompt_tokens: int = 10,
    completion_tokens: int = 5,
    cost_usd: float = 0.001,
    latency_ms: int = 25,
) -> LlmUsageRecord:
    """One usage row with controllable numbers, for aggregate-read tests."""
    return LlmUsageRecord(
        tenant_id=tenant_id,
        job_type=job_type,
        job_id=job_id,
        model="gpt-4o-mini",
        prompt_tokens=prompt_tokens,
        completion_tokens=completion_tokens,
        cost_usd=cost_usd,
        latency_ms=latency_ms,
        request_id="",
    )


class TestUsageRepository:
    async def test_record_persists_one_row_per_call_with_its_facts(self) -> None:
        tenant_id = uuid.uuid4()
        usage_record = LlmUsage(
            model="gpt-4o-mini",
            prompt_tokens=120,
            completion_tokens=30,
            cost_usd=0.000036,
            latency_ms=812,
        )

        async with session() as s:
            await usage_repo.record(
                s,
                tenant_id=tenant_id,
                job_type=UsageJobKind.extraction,
                job_id=uuid.uuid4(),
                usage=usage_record,
                request_id="req-1",
            )

        async with session() as s:
            rows = (
                (
                    await s.execute(
                        select(LlmUsageRecord).where(LlmUsageRecord.tenant_id == tenant_id)
                    )
                )
                .scalars()
                .all()
            )
        assert len(rows) == 1
        row = rows[0]
        assert row.tenant_id == tenant_id
        assert row.job_type is UsageJobKind.extraction
        assert row.model == "gpt-4o-mini"
        assert row.prompt_tokens == 120
        assert row.completion_tokens == 30
        assert row.cost_usd == pytest.approx(0.000036)
        assert row.latency_ms == 812
        assert row.request_id == "req-1"
        assert row.created_at is not None

    async def test_spend_by_job_type_sums_per_kind_and_stays_tenant_scoped(self) -> None:
        mine, theirs = uuid.uuid4(), uuid.uuid4()
        async with session() as s:
            s.add_all(
                [
                    usage(tenant_id=mine, job_type=UsageJobKind.extraction, cost_usd=0.010),
                    usage(tenant_id=mine, job_type=UsageJobKind.extraction, cost_usd=0.002),
                    usage(tenant_id=mine, job_type=UsageJobKind.ingestion, cost_usd=0.001),
                    usage(tenant_id=theirs, job_type=UsageJobKind.extraction, cost_usd=9.0),
                ]
            )
            await s.commit()

        async with session() as s:
            rows = await usage_repo.spend_by_job_type(s, tenant_id=mine)

        assert [(row.job_type, row.calls) for row in rows] == [
            ("ingestion", 1),
            ("extraction", 2),
        ]
        extraction = rows[1]
        assert extraction.job_id is None
        assert extraction.prompt_tokens == 20
        assert extraction.completion_tokens == 10
        assert extraction.cost_usd == pytest.approx(0.012)
        assert extraction.latency_ms == 50

    async def test_spend_by_job_groups_calls_by_their_job_id(self) -> None:
        tenant_id = uuid.uuid4()
        job_a, job_b = uuid.uuid4(), uuid.uuid4()
        async with session() as s:
            s.add_all(
                [
                    usage(tenant_id=tenant_id, job_type=UsageJobKind.extraction, job_id=job_a),
                    usage(tenant_id=tenant_id, job_type=UsageJobKind.extraction, job_id=job_a),
                    usage(tenant_id=tenant_id, job_type=UsageJobKind.ingestion, job_id=job_b),
                    # No job row (the search request path): its own group.
                    usage(tenant_id=tenant_id, job_type=UsageJobKind.search, job_id=None),
                ]
            )
            await s.commit()

        async with session() as s:
            rows = await usage_repo.spend_by_job(s, tenant_id=tenant_id)

        assert [(row.job_type, row.job_id, row.calls) for row in rows] == [
            ("ingestion", job_b, 1),
            ("extraction", job_a, 2),
            ("search", None, 1),
        ]

    async def test_spend_total_sums_every_row_for_the_tenant(self) -> None:
        mine, theirs = uuid.uuid4(), uuid.uuid4()
        async with session() as s:
            s.add_all(
                [
                    usage(tenant_id=mine, job_type=UsageJobKind.extraction),
                    usage(tenant_id=mine, job_type=UsageJobKind.search, job_id=None),
                    usage(tenant_id=theirs, job_type=UsageJobKind.extraction),
                ]
            )
            await s.commit()

        async with session() as s:
            total = await usage_repo.spend_total(s, tenant_id=mine)

        assert total.calls == 2
        assert total.prompt_tokens == 20
        assert total.completion_tokens == 10
        assert total.cost_usd == pytest.approx(0.002)
        assert total.latency_ms == 50


class TestExtractionPersistence:
    async def test_run_persists_one_row_for_the_extraction_call(self) -> None:
        tenant_id = uuid.uuid4()
        document = await make_text_document(tenant_id, text="0123456789")
        async with session() as s:
            job = await enqueue_extraction(s, tenant_id=tenant_id, document_id=document.id)

        async with (
            fake_llm_client(EXTRACTION_OUTPUT, prompt_tokens=50, completion_tokens=10) as (
                llm,
                _,
            ),
            session() as s,
        ):
            await run_next_extraction_job(s, llm=llm, review_threshold=0.8)

        rows = await tenant_usage_rows(tenant_id)
        assert len(rows) == 1
        row = rows[0]
        assert row.job_type is UsageJobKind.extraction
        assert row.job_id == job.id
        assert row.model == "gpt-4o-mini"
        assert row.prompt_tokens == 50
        assert row.completion_tokens == 10
        assert row.cost_usd > 0.0
        assert row.latency_ms >= 0

    async def test_citation_gate_retry_persists_one_row_per_attempt(self) -> None:
        """Each attempt is its own LLM call with its own log line: two calls,
        two usage rows — the spend of the rejected attempt is still accounted."""
        tenant_id = uuid.uuid4()
        document = await make_text_document(tenant_id, text="0123456789")
        async with session() as s:
            await enqueue_extraction(s, tenant_id=tenant_id, document_id=document.id)
        # char_end 10_000 is far outside the 10-char text: an invalid span.
        invalid = EXTRACTION_OUTPUT.replace('"char_end": 8', '"char_end": 10000')

        async with (
            fake_llm_client_queue([invalid, EXTRACTION_OUTPUT]) as (llm, _),
            session() as s,
        ):
            await run_next_extraction_job(s, llm=llm, review_threshold=0.8)

        assert len(await tenant_usage_rows(tenant_id)) == 2, "one row per call, retries included"


class TestIngestionPersistence:
    async def test_run_persists_one_row_per_embed_batch(self, data_dir: Path) -> None:
        tenant_id = uuid.uuid4()
        document = await make_docx_document(tenant_id, data_dir=data_dir)
        async with session() as s:
            job = await enqueue_ingestion(s, tenant_id=tenant_id, document_id=document.id)
        vectors = [[0.5] * EMBEDDING_DIMENSIONS, [0.25] * EMBEDDING_DIMENSIONS]

        async with (
            fake_embedding_client(vectors, prompt_tokens=120) as (llm, _),
            session() as s,
        ):
            await run_next_ingestion_job(s, llm=llm, data_dir=str(data_dir))

        rows = await tenant_usage_rows(tenant_id)
        assert len(rows) == 1
        row = rows[0]
        assert row.job_type is UsageJobKind.ingestion
        assert row.job_id == job.id
        assert row.model == "text-embedding-3-small"
        assert row.prompt_tokens == 120
        assert row.completion_tokens == 0
        assert row.latency_ms >= 0


class TestChangeReportPersistence:
    async def test_run_persists_rows_for_every_call_under_the_report_job(self) -> None:
        base, amended = await ready_pair_with_base_obligation()
        async with session() as s:
            job = await enqueue_change_report(
                s,
                tenant_id=base.tenant_id,
                base_document_id=base.id,
                amended_document_id=amended.id,
            )

        async with (
            fake_llm_with_embeddings_client(
                [EXPLANATION_OUTPUT, IMPACT_OUTPUT], [CHUNK_EMBEDDING]
            ) as (llm, requests),
            session() as s,
        ):
            await run_next_change_report_job(s, llm=llm, review_threshold=0.8)

        assert len(requests) == 3, "one explanation, one embed, one impact call"
        rows = await tenant_usage_rows(base.tenant_id)
        assert len(rows) == 3
        for row in rows:
            assert row.job_type is UsageJobKind.change_report
            assert row.job_id == job.id
            assert row.latency_ms >= 0
        assert {row.model for row in rows if row.completion_tokens > 0} == {"gpt-4o-mini"}
        assert any(row.model == "text-embedding-3-small" for row in rows)


class TestSearchPersistence:
    async def test_search_request_persists_a_row_under_the_search_surface(
        self, client: AsyncClient
    ) -> None:
        """GET /search has no job row: its query embed lands under the search
        kind with job_id None, correlated to the request id."""
        tenant_id = uuid.uuid4()
        await seed_searchable_chunk(tenant_id)
        override_llm_with_query_vector([1.0] + [0.0] * (EMBEDDING_DIMENSIONS - 1))

        response = await client.get("/search", params={"q": "payment"}, headers=bearer(tenant_id))

        assert response.status_code == 200
        rows = await tenant_usage_rows(tenant_id)
        assert len(rows) == 1
        row = rows[0]
        assert row.job_type is UsageJobKind.search
        assert row.job_id is None
        assert row.model == "text-embedding-3-small"
        assert row.prompt_tokens == 1
        assert row.completion_tokens == 0
        assert row.request_id == response.headers["x-request-id"]


# --- shared helpers ---------------------------------------------------------


async def seed_searchable_chunk(tenant_id: uuid.UUID) -> None:
    """One document with one chunk whose embedding matches itself."""
    text = "Supplier shall pay within thirty days."
    async with session() as s:
        document = await documents_repo.create(
            s,
            tenant_id=tenant_id,
            filename="msa.txt",
            mime_type="text/plain",
            sha256=uuid.uuid4().hex * 2,
        )
        await texts_repo.replace(
            s, tenant_id=tenant_id, document_id=document.id, text=text, page_map=[]
        )
        await chunks_repo.replace_with_embeddings(
            s,
            tenant_id=tenant_id,
            document_id=document.id,
            chunks=[(text, 0, len(text))],
            embeddings=[[1.0] + [0.0] * (EMBEDDING_DIMENSIONS - 1)],
        )
        await s.commit()


def override_llm_with_query_vector(vector: list[float]) -> None:
    async def override() -> AsyncIterator[OpenAiClient]:
        async with fake_embedding_client([vector], prompt_tokens=1) as (llm, _):
            yield llm

    app.dependency_overrides[get_llm_client] = override


async def tenant_usage_rows(tenant_id: uuid.UUID) -> list[LlmUsageRecord]:
    async with session() as s:
        return list(
            (await s.execute(select(LlmUsageRecord).where(LlmUsageRecord.tenant_id == tenant_id)))
            .scalars()
            .all()
        )


async def make_text_document(tenant_id: uuid.UUID, *, text: str) -> Document:
    """One parsed text document ready for extraction."""
    async with session() as s:
        document = await documents_repo.create(
            s,
            tenant_id=tenant_id,
            filename="agreement.txt",
            mime_type="text/plain",
            sha256=uuid.uuid4().hex * 2,
        )
        await texts_repo.replace(
            s, tenant_id=tenant_id, document_id=document.id, text=text, page_map=[]
        )
        await documents_repo.set_status(s, document, DocumentStatus.parsed)
        return document


async def make_docx_document(tenant_id: uuid.UUID, *, data_dir: Path) -> Document:
    async with session() as s:
        document = await documents_repo.create(
            s,
            tenant_id=tenant_id,
            filename="msa.docx",
            mime_type=DOCX_MIME,
            sha256=uuid.uuid4().hex * 2,
        )
    local.save_document(str(data_dir), tenant_id, document.sha256, contract_docx())
    return document


async def make_document(
    *, tenant_id: uuid.UUID, text: str, parsed: bool, amends: Document | None = None
) -> Document:
    async with session() as s:
        document = await documents_repo.create(
            s,
            tenant_id=tenant_id,
            filename="agreement.txt",
            mime_type="text/plain",
            sha256=uuid.uuid4().hex * 2,
            amends_document_id=amends.id if amends else None,
        )
        await texts_repo.replace(
            s, tenant_id=tenant_id, document_id=document.id, text=text, page_map=[]
        )
        if parsed:
            await documents_repo.set_status(s, document, DocumentStatus.parsed)
        return document


async def complete_extraction_with_delivery_obligation(document: Document) -> None:
    cite_start = BASE_TEXT.index("LICENSOR shall deliver")
    async with session() as s:
        s.add(
            ExtractionJob(
                tenant_id=document.tenant_id,
                document_id=document.id,
                status=ExtractionJobStatus.completed,
                result={
                    "obligations": [
                        {
                            "clause_ref": "2.1",
                            "description": "Licensor shall deliver within 14 days.",
                            "owner": "Licensor",
                            "citation": {
                                "char_start": cite_start,
                                "char_end": len(BASE_TEXT),
                            },
                            "confidence": 0.9,
                        }
                    ],
                    "defined_terms": [],
                },
            )
        )
        s.add(
            DocumentChunk(
                tenant_id=document.tenant_id,
                document_id=document.id,
                ordinal=0,
                text=BASE_TEXT,
                char_start=0,
                char_end=len(BASE_TEXT),
                embedding=CHUNK_EMBEDDING,
            )
        )
        await s.commit()


async def ready_pair_with_base_obligation() -> tuple[Document, Document]:
    tenant_id = uuid.uuid4()
    base = await make_document(tenant_id=tenant_id, text=BASE_TEXT, parsed=True)
    amended = await make_document(tenant_id=tenant_id, text=AMENDED_TEXT, parsed=True, amends=base)
    await complete_extraction_with_delivery_obligation(base)
    async with session() as s:
        s.add(
            ExtractionJob(
                tenant_id=tenant_id,
                document_id=amended.id,
                status=ExtractionJobStatus.completed,
                result={"obligations": [], "defined_terms": []},
            )
        )
        await s.commit()
    return base, amended
