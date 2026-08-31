"""Integration tests for the ingestion pipeline service: enqueue, run, re-run.

Real Postgres (job/document/text/chunk rows); the embedding wire is faked at
the httpx2 transport seam via tests/fake_openai, per the AGENTS.md exception
for LLM tests.
"""

import io
import json
import uuid
from collections.abc import AsyncIterator, Iterator
from contextlib import asynccontextmanager
from datetime import UTC, datetime, timedelta
from pathlib import Path

import httpx2
import pytest
from docx import Document as DocxDocument
from sqlalchemy import delete, func, select
from sqlalchemy.ext.asyncio import AsyncSession

import app.db as db
from app.config import get_settings
from app.llm.client import OpenAiClient
from app.models.document import Document, DocumentStatus
from app.models.document_chunk import EMBEDDING_DIMENSIONS, DocumentChunk
from app.models.document_text import DocumentText
from app.models.ingestion_job import IngestionJob, IngestionJobStatus
from app.repositories import document_chunks as chunks_repo
from app.repositories import document_texts as texts_repo
from app.repositories import documents as documents_repo
from app.repositories import ingestion_jobs as jobs_repo
from app.services.documents import DocumentNotFoundError
from app.services.ingestion import (
    IngestionJobConflictError,
    IngestionJobNotFoundError,
    enqueue_ingestion,
    get_ingestion_job,
    run_next_ingestion_job,
)
from app.storage import local
from tests.fake_openai import fake_embedding_client, make_settings

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def contract_docx() -> bytes:
    """A two-clause DOCX; each clause becomes one chunk (ADR-005)."""
    buffer = io.BytesIO()
    doc = DocxDocument()
    doc.add_heading("Section 1: Payment", level=1)
    doc.add_paragraph("Supplier shall pay within thirty days of invoice.")
    doc.add_heading("Section 2: Delivery", level=1)
    doc.add_paragraph("Goods shall be delivered to the customer site.")
    doc.save(buffer)
    return buffer.getvalue()


def fake_vector(seed: int) -> list[float]:
    return [0.001 * ((seed + i) % 10) for i in range(EMBEDDING_DIMENSIONS)]


@asynccontextmanager
async def session() -> AsyncIterator[AsyncSession]:
    async with db.get_sessionmaker()() as s:
        yield s


@pytest.fixture(autouse=True)
async def engine() -> AsyncIterator[None]:
    db.init_engine(get_settings().database_url)
    yield
    await db.dispose_engine()


@pytest.fixture(autouse=True)
def data_dir(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> Iterator[Path]:
    monkeypatch.setenv("DATA_DIR", str(tmp_path))
    yield tmp_path


@pytest.fixture(autouse=True)
async def clean_tables() -> AsyncIterator[None]:
    yield
    async with session() as s:
        await s.execute(delete(DocumentChunk))
        await s.execute(delete(DocumentText))
        await s.execute(delete(IngestionJob))
        await s.execute(delete(Document))
        await s.commit()


async def make_document(
    tenant_id: uuid.UUID,
    *,
    mime_type: str = DOCX_MIME,
    content: bytes | None = None,
    data_dir: Path,
) -> Document:
    async with session() as s:
        document = await documents_repo.create(
            s,
            tenant_id=tenant_id,
            filename="msa.docx",
            mime_type=mime_type,
            sha256=uuid.uuid4().hex * 2,
        )
    payload = content if content is not None else contract_docx()
    local.save_document(str(data_dir), tenant_id, document.sha256, payload)
    return document


async def enqueue(tenant_id: uuid.UUID, document_id: uuid.UUID) -> IngestionJob:
    async with session() as s:
        return await enqueue_ingestion(s, tenant_id=tenant_id, document_id=document_id)


class TestEnqueueIngestion:
    async def test_enqueue_creates_queued_job_for_existing_document(self, data_dir: Path) -> None:
        tenant_id = uuid.uuid4()
        document = await make_document(tenant_id, data_dir=data_dir)

        job = await enqueue(tenant_id, document.id)

        assert job.status is IngestionJobStatus.queued
        assert job.tenant_id == tenant_id
        assert job.document_id == document.id
        assert job.result is None
        assert job.error is None

    async def test_enqueue_unknown_document_raises_document_not_found(self) -> None:
        with pytest.raises(DocumentNotFoundError):
            async with session() as s:
                await enqueue_ingestion(s, tenant_id=uuid.uuid4(), document_id=uuid.uuid4())

    async def test_enqueue_while_a_job_is_active_conflicts_with_existing_job_id(
        self, data_dir: Path
    ) -> None:
        tenant_id = uuid.uuid4()
        document = await make_document(tenant_id, data_dir=data_dir)
        first = await enqueue(tenant_id, document.id)

        with pytest.raises(IngestionJobConflictError) as excinfo:
            await enqueue(tenant_id, document.id)

        assert excinfo.value.job_id == first.id

    async def test_enqueue_after_a_terminal_job_starts_a_new_run(self, data_dir: Path) -> None:
        tenant_id = uuid.uuid4()
        document = await make_document(tenant_id, data_dir=data_dir)
        first = await enqueue(tenant_id, document.id)
        async with session() as s:
            stored = await s.get(IngestionJob, first.id)
            assert stored is not None
            stored.status = IngestionJobStatus.completed
            await s.commit()

        second = await enqueue(tenant_id, document.id)

        assert second.id != first.id
        assert second.status is IngestionJobStatus.queued


class TestGetIngestionJob:
    async def test_get_returns_job_scoped_to_tenant(self, data_dir: Path) -> None:
        tenant_id = uuid.uuid4()
        document = await make_document(tenant_id, data_dir=data_dir)
        job = await enqueue(tenant_id, document.id)

        async with session() as s:
            found = await get_ingestion_job(s, tenant_id=tenant_id, job_id=job.id)

        assert found.id == job.id

    async def test_other_tenant_or_unknown_job_reads_as_not_found(self, data_dir: Path) -> None:
        tenant_id = uuid.uuid4()
        document = await make_document(tenant_id, data_dir=data_dir)
        job = await enqueue(tenant_id, document.id)

        with pytest.raises(IngestionJobNotFoundError):
            async with session() as s:
                await get_ingestion_job(s, tenant_id=uuid.uuid4(), job_id=job.id)
        with pytest.raises(IngestionJobNotFoundError):
            async with session() as s:
                await get_ingestion_job(s, tenant_id=tenant_id, job_id=uuid.uuid4())


class TestRunNextIngestionJob:
    async def test_no_queued_job_returns_false(self) -> None:
        async with fake_embedding_client([fake_vector(0)]) as (llm, _), session() as s:
            ran = await run_next_ingestion_job(s, llm=llm, data_dir="/nonexistent")

        assert ran is False

    async def test_docx_completes_with_text_clauses_and_embeddings(self, data_dir: Path) -> None:
        tenant_id = uuid.uuid4()
        document = await make_document(tenant_id, data_dir=data_dir)
        job = await enqueue(tenant_id, document.id)
        vectors = [fake_vector(1), fake_vector(2)]

        async with (
            fake_embedding_client(vectors, prompt_tokens=120) as (llm, requests),
            session() as s,
        ):
            ran = await run_next_ingestion_job(s, llm=llm, data_dir=str(data_dir))

        assert ran is True
        async with session() as s:
            completed = await jobs_repo.find_by_id(s, tenant_id=tenant_id, job_id=job.id)
            text_row = await texts_repo.find_by_document_id(s, document_id=document.id)
            chunk_rows = (
                (
                    await s.execute(
                        select(DocumentChunk)
                        .where(DocumentChunk.document_id == document.id)
                        .order_by(DocumentChunk.ordinal)
                    )
                )
                .scalars()
                .all()
            )
            doc = await documents_repo.find_by_id(s, tenant_id=tenant_id, document_id=document.id)
        assert completed is not None
        assert completed.status is IngestionJobStatus.completed
        assert completed.result == {"chunk_count": 2}
        assert completed.error is None
        assert doc is not None
        assert doc.status is DocumentStatus.parsed
        assert text_row is not None
        # Every chunk cites itself: text equals its span of the parsed text.
        assert len(chunk_rows) == 2
        for ordinal, chunk_row in enumerate(chunk_rows):
            assert chunk_row.ordinal == ordinal
            assert text_row.text[chunk_row.char_start : chunk_row.char_end] == chunk_row.text
        # pgvector stores float4, so compare numerically, not bit-exact.
        for chunk_row, vector in zip(chunk_rows, vectors, strict=True):
            assert list(chunk_row.embedding) == pytest.approx(vector)
        # The chunk texts travelled to the embeddings endpoint, in order.
        sent = json.loads(requests[0].content)
        assert sent["input"] == [row.text for row in chunk_rows]
        assert sent["model"] == "text-embedding-3-small"

    async def test_empty_document_completes_with_zero_chunks_and_no_llm_call(
        self, data_dir: Path
    ) -> None:
        tenant_id = uuid.uuid4()
        document = await make_document(
            tenant_id, mime_type="text/plain", content=b"", data_dir=data_dir
        )
        job = await enqueue(tenant_id, document.id)

        async with fake_embedding_client([fake_vector(0)]) as (llm, requests), session() as s:
            ran = await run_next_ingestion_job(s, llm=llm, data_dir=str(data_dir))

        assert ran is True
        assert requests == [], "an empty document must not spend embedding tokens"
        async with session() as s:
            completed = await jobs_repo.find_by_id(s, tenant_id=tenant_id, job_id=job.id)
            doc = await documents_repo.find_by_id(s, tenant_id=tenant_id, document_id=document.id)
        assert completed is not None
        assert completed.status is IngestionJobStatus.completed
        assert completed.result == {"chunk_count": 0}
        assert doc is not None
        assert doc.status is DocumentStatus.parsed

    async def test_unparseable_bytes_fail_the_job_and_flag_the_document(
        self, data_dir: Path
    ) -> None:
        tenant_id = uuid.uuid4()
        document = await make_document(
            tenant_id,
            mime_type="application/pdf",
            content=b"definitely not a pdf",
            data_dir=data_dir,
        )
        job = await enqueue(tenant_id, document.id)

        async with fake_embedding_client([fake_vector(0)]) as (llm, requests), session() as s:
            ran = await run_next_ingestion_job(s, llm=llm, data_dir=str(data_dir))

        assert ran is True, "the handler must swallow parse failures so the loop keeps going"
        assert requests == [], "no embedding call may happen when parsing already failed"
        async with session() as s:
            failed = await jobs_repo.find_by_id(s, tenant_id=tenant_id, job_id=job.id)
            doc = await documents_repo.find_by_id(s, tenant_id=tenant_id, document_id=document.id)
            text_count = (
                await s.execute(select(func.count()).select_from(DocumentText))
            ).scalar_one()
        assert failed is not None
        assert failed.status is IngestionJobStatus.failed
        assert failed.result is None
        assert failed.error, "the failure reason must be recorded for the poller"
        assert doc is not None
        assert doc.status is DocumentStatus.failed
        assert text_count == 0, "no parsed text may be left behind by a failed run"

    async def test_embedding_wire_failure_fails_the_job_without_partial_writes(
        self, data_dir: Path
    ) -> None:
        tenant_id = uuid.uuid4()
        document = await make_document(tenant_id, data_dir=data_dir)
        job = await enqueue(tenant_id, document.id)

        # A wire that always 500s, so embeddings fail after parsing succeeded.
        wire = httpx2.AsyncClient(
            transport=httpx2.MockTransport(lambda request: httpx2.Response(500, content=b"boom"))
        )
        llm = OpenAiClient(make_settings(), http_client=wire)
        try:
            async with session() as s:
                ran = await run_next_ingestion_job(s, llm=llm, data_dir=str(data_dir))
        finally:
            await wire.aclose()

        assert ran is True
        async with session() as s:
            failed = await jobs_repo.find_by_id(s, tenant_id=tenant_id, job_id=job.id)
            doc = await documents_repo.find_by_id(s, tenant_id=tenant_id, document_id=document.id)
            chunk_count = (
                await s.execute(select(func.count()).select_from(DocumentChunk))
            ).scalar_one()
        assert failed is not None
        assert failed.status is IngestionJobStatus.failed
        assert "llm call failed" in failed.error
        assert doc is not None
        assert doc.status is DocumentStatus.failed
        assert chunk_count == 0, "a failed embedding run must not persist half the pipeline"

    async def test_attempt_capped_claim_fails_the_document_without_running(
        self, data_dir: Path
    ) -> None:
        # A job whose lease kept lapsing (crashed worker) past the attempt cap
        # claims into failed; the document must flip to failed, not stay
        # `uploaded` forever, and no embed work may happen.
        tenant_id = uuid.uuid4()
        document = await make_document(tenant_id, data_dir=data_dir)
        job = await enqueue(tenant_id, document.id)
        async with session() as s:
            stored = await s.get(IngestionJob, job.id)
            assert stored is not None
            stored.status = IngestionJobStatus.running
            stored.attempts = 3
            stored.lease_expires_at = datetime.now(UTC) - timedelta(seconds=1)
            await s.commit()

        async with fake_embedding_client([fake_vector(0)]) as (llm, requests), session() as s:
            ran = await run_next_ingestion_job(s, llm=llm, data_dir=str(data_dir))

        assert ran is True
        assert requests == [], "a capped job must not run any embed work"
        async with session() as s:
            failed = await jobs_repo.find_by_id(s, tenant_id=tenant_id, job_id=job.id)
            doc = await documents_repo.find_by_id(s, tenant_id=tenant_id, document_id=document.id)
        assert failed is not None
        assert failed.status is IngestionJobStatus.failed
        assert failed.error == "max_attempts_exceeded"
        assert doc is not None
        assert doc.status is DocumentStatus.failed

    async def test_rerun_after_completion_replaces_text_and_chunks(self, data_dir: Path) -> None:
        tenant_id = uuid.uuid4()
        document = await make_document(tenant_id, data_dir=data_dir)
        vectors = [fake_vector(1), fake_vector(2)]
        await enqueue(tenant_id, document.id)
        async with fake_embedding_client(vectors) as (llm, _), session() as s:
            await run_next_ingestion_job(s, llm=llm, data_dir=str(data_dir))

        second = await enqueue(tenant_id, document.id)
        async with fake_embedding_client(vectors) as (llm, _), session() as s:
            ran = await run_next_ingestion_job(s, llm=llm, data_dir=str(data_dir))

        assert ran is True
        async with session() as s:
            completed = await jobs_repo.find_by_id(s, tenant_id=tenant_id, job_id=second.id)
            chunk_count = await chunks_repo.count_for_document(s, document_id=document.id)
            text_rows = (
                await s.execute(select(func.count()).select_from(DocumentText))
            ).scalar_one()
        assert completed is not None
        assert completed.status is IngestionJobStatus.completed
        assert chunk_count == 2, "a re-run replaces chunks, never accumulates them"
        assert text_rows == 1, "a re-run replaces the parsed-text row, never duplicates it"
